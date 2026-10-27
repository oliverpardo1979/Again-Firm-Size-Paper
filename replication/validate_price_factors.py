"""Check the adopted deflators against DANE's national December CPI series."""
import argparse
import hashlib
import json
from pathlib import Path
import pandas as pd
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser()
parser.add_argument('workbook', type=Path)
parser.add_argument('--source-url', required=True)
parser.add_argument('--output', type=Path, required=True)
args = parser.parse_args()
sheet = openpyxl.load_workbook(args.workbook, data_only=True)['IndicesIPC']
header = next(row for row in sheet.iter_rows() if row[0].value == 'Mes')
december = next(row for row in sheet.iter_rows() if str(row[0].value).strip().lower() == 'diciembre')
levels = {int(h.value): float(v.value) for h, v in zip(header[1:], december[1:])
          if isinstance(h.value, (int, float)) and isinstance(v.value, (int, float))}
adopted = json.loads((ROOT / 'replication/primary_income_output/price_factors.json').read_text())
rows = []
for year, factor in adopted.items():
    calculated = levels[2025] / levels[int(year)]
    error = abs(calculated - factor)
    assert error < 1e-12, (year, calculated, factor)
    rows.append({'year': int(year), 'december_cpi_2018_100': levels[int(year)],
                 'price_factor_2025': factor, 'absolute_error': error})
args.output.mkdir(parents=True, exist_ok=True)
pd.DataFrame(rows).to_csv(args.output / 'cpi_validation.csv', index=False)
(args.output / 'cpi_source.json').write_text(json.dumps({
    'source': 'DANE, national CPI linked index, December 2018 = 100',
    'url': args.source_url, 'sheet': 'IndicesIPC', 'month': 'December',
    'workbook_sha256': hashlib.sha256(args.workbook.read_bytes()).hexdigest(),
    'factors_checked': len(rows), 'maximum_absolute_error': max(r['absolute_error'] for r in rows),
}, indent=2) + '\n')
print(f'Validated {len(rows)} annual deflators against the DANE workbook.')
